"""Regenerate the group presentation template.

    pip install python-pptx
    python scripts/make_deck_template.py

Built on the group's earlier deck - white ground, title top left over a thin
rule, institution mark top right, page number bottom right, check bullets - with
the MS3L mark added beside KRICT's and the rule redrawn in the KRICT CI ramp.

Writes into assets/templates/. Change a layout or a colour here and re-run; do
not edit the slides by hand, or the deck, the poster and the website stop
agreeing.
"""
import os
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.makedirs("assets/templates", exist_ok=True)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

FONT = "Arial"                      # what the group's decks already use
NAVY = RGBColor(0x0B, 0x2F, 0x5B)
BLUE = RGBColor(0x00, 0x75, 0xC2)   # KRICT PANTONE 3015C
TEAL = RGBColor(0x00, 0xAD, 0xA9)   # KRICT PANTONE 632C
GREY = RGBColor(0x5B, 0x6F, 0x84)
LINE = RGBColor(0xD6, 0xE2, 0xEE)
LIGHT = RGBColor(0xF4, 0xF8, 0xFD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
L = "assets/images/logo"
LOCK = f"{L}/ms3l-lockup-horizontal.png"
MARK = f"{L}/ms3l-avatar-circle.png"
KRICT = "assets/images/logos/krict-logo.png"

M = Inches(0.62)                    # left margin, matching the old deck


def grad(sh, angle=0.0):
    sh.fill.gradient()
    st = sh.fill.gradient_stops
    st[0].color.rgb, st[0].position = BLUE, 0.0
    st[1].color.rgb, st[1].position = TEAL, 1.0
    sh.fill.gradient_angle = angle
    sh.line.fill.background()


def rect(s, x, y, w, h):
    return s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)


def solid(sh, c, line=None):
    sh.fill.solid(); sh.fill.fore_color.rgb = c
    if line is not None:
        sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()


def text(s, x, y, w, h, t, size, color, bold=False, align=PP_ALIGN.LEFT, space=1.0):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = space
    r = p.add_run(); r.text = t
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = color
    return tb


def chrome(s, n, title=None):
    """Marks top right, page number bottom right, and the header rule."""
    s.shapes.add_picture(MARK, W - Inches(1.72), Inches(0.24), height=Inches(0.42))
    solid(rect(s, W - Inches(1.19), Inches(0.28), Inches(0.01), Inches(0.34)), LINE)
    s.shapes.add_picture(KRICT, W - Inches(1.10), Inches(0.32), height=Inches(0.26))
    if title is not None:
        text(s, M, Inches(0.20), Inches(9.4), Inches(0.6), title, 28, BLUE, True)
    grad(rect(s, M, Inches(0.82), W - 2 * M, Inches(0.035)))
    text(s, W - Inches(1.30), H - Inches(0.52), Inches(0.7), Inches(0.3),
         str(n), 12, GREY, align=PP_ALIGN.RIGHT)


def bullets(s, items, top=Inches(1.58), size=15, gap=0.66):
    for i, (head, sub) in enumerate(items):
        y = top + Inches(i * gap)
        text(s, M + Inches(0.06), y, Inches(0.32), Inches(0.34), "✓", size, TEAL, True)
        text(s, M + Inches(0.40), y, Inches(11.6), Inches(0.34), head, size, NAVY, True)
        if sub:
            text(s, M + Inches(0.40), y + Inches(0.30), Inches(11.6), Inches(0.32),
                 sub, size - 2, GREY)


# 1 -- title ----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(LOCK, M, Inches(0.55), width=Inches(4.6))
s.shapes.add_picture(KRICT, W - Inches(2.10), Inches(0.62), height=Inches(0.40))
grad(rect(s, M, Inches(1.62), W - 2 * M, Inches(0.045)))
text(s, M, Inches(2.62), W - 2 * M, Inches(1.3), "Title", 46, NAVY, True, PP_ALIGN.CENTER)
text(s, M, Inches(3.92), W - 2 * M, Inches(0.7), "Sub-title", 26, BLUE, False, PP_ALIGN.CENTER)
grad(rect(s, Inches(5.92), Inches(4.86), Inches(1.5), Inches(0.045)))
text(s, M, Inches(5.42), W - 2 * M, Inches(0.5), "Name", 18, NAVY, True, PP_ALIGN.CENTER)
text(s, M, Inches(5.86), W - 2 * M, Inches(0.5),
     "Membrane-based Sustainable Separation Solutions Laboratory, KRICT", 14, GREY,
     False, PP_ALIGN.CENTER)
text(s, M, Inches(6.24), W - 2 * M, Inches(0.4), "2026.00.00.", 13, GREY, False, PP_ALIGN.CENTER)

# 2 -- outline --------------------------------------------------------------
s = prs.slides.add_slide(BLANK); chrome(s, 2, "Outline")
bullets(s, [("Content #1", None), ("Content #2", None), ("Content #3", None)],
        top=Inches(1.40), size=18, gap=0.78)

# 3 -- content --------------------------------------------------------------
s = prs.slides.add_slide(BLANK); chrome(s, 3, "Content #1")
text(s, M, Inches(1.02), Inches(11.6), Inches(0.4), "Subjective", 15, BLUE, True)
bullets(s, [("Point one", "Supporting detail for the first point."),
            ("Point two", "Supporting detail for the second point."),
            ("Point three", "Supporting detail for the third point.")])

# 4 -- two column -----------------------------------------------------------
s = prs.slides.add_slide(BLANK); chrome(s, 4, "Content #2")
text(s, M, Inches(1.02), Inches(11.6), Inches(0.4), "Subjective", 15, BLUE, True)
for i, side in enumerate(["Left column", "Right column"]):
    x = M + Inches(i * 6.14)
    solid(rect(s, x, Inches(1.52), Inches(5.85), Inches(4.9)), LIGHT, LINE)
    grad(rect(s, x, Inches(1.52), Inches(0.05), Inches(4.9)), 270.0)
    text(s, x + Inches(0.34), Inches(1.82), Inches(5.2), Inches(0.4), side, 17, NAVY, True)
    text(s, x + Inches(0.34), Inches(2.30), Inches(5.2), Inches(3.6),
         "Body text. Replace with a figure, a table or bullets.", 14, GREY, space=1.35)

# 5 -- figure ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK); chrome(s, 5, "Content #3")
ph = rect(s, M, Inches(1.30), W - 2 * M, Inches(5.0)); solid(ph, LIGHT, LINE)
text(s, M, Inches(3.62), W - 2 * M, Inches(0.5), "Place figure here", 15,
     RGBColor(0x9A, 0xAE, 0xC2), align=PP_ALIGN.CENTER)
text(s, M, Inches(6.44), W - 2 * M, Inches(0.4), "Figure 1. Caption.", 12, GREY)

# 6 -- conclusions ----------------------------------------------------------
s = prs.slides.add_slide(BLANK); chrome(s, 6, "Conclusions")
bullets(s, [("Summary point one", None), ("Summary point two", None),
            ("Summary point three", None)], top=Inches(1.40), size=18, gap=0.78)
text(s, M, Inches(5.62), Inches(11.6), Inches(0.4),
     "jh.kim@krict.re.kr    |    https://ms3l.org", 14, BLUE)

prs.save("assets/templates/MS3L_presentation_template.pptx")
print("deck rebuilt")
